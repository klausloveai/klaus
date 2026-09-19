#!/usr/bin/env python3
"""Lingtu Law — Case Referral "Lead Details" builder.

Question架构 came from Morgan & Morgan (Babatunde Oyelowo, 04/27/2026).
Field names are kept VERBATIM; only the formatting is ours.

Usage:  python3 build_referral_lead_details.py <out.docx> [data.json]
No data.json -> emits the blank TEMPLATE with [bracketed] placeholders.
"""
import json, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
INK  = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x60, 0x60, 0x60)
RULE = "808080"

# (section heading, [(label, key, indented?), ...]) -- order is Bob's original order.
SCHEMA = [
    ("Lead", [
        ("Name",                         "name",        False),
        ("Case Type",                    "case_type",   False),
    ]),
    ("Incident", [
        ("Where did the accident happen?", "where",     False),
        ("Incident Date",                "date",        False),
        ("Summary",                      "summary",     False),
        ("Police on scene",              "police",      False),
    ]),
    ("Coverage", [
        ("PC's Insurance",               "pc_ins",      False),
        ("Limits",                       "pc_limits",   True),
        ("UIM/UM",                       "um",          False),
        ("Limits",                       "um_limits",   True),
        ("Def's Insurance",              "def_ins",     False),
        ("Limits",                       "def_limits",  True),
    ]),
    ("Injuries & Treatment", [
        ("Injuries",                     "injuries",    False),
        ("Emergency room or Urgent Care","er",          False),
        ("Transported via Ambulance",    "ambulance",   False),
        ("Health Insurance",             "health_ins",  False),
        ("Treatment",                    "treatment",   False),
    ]),
]

BLANK = {k: f"[{lbl}]" for _, fields in SCHEMA for lbl, k, _ in fields}


def _rule(p, size=6):
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), str(size))
    bot.set(qn('w:space'), '2');    bot.set(qn('w:color'), RULE)
    bd.append(bot); pr.append(bd)


def build(data, out):
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin  = s.bottom_margin = Inches(0.8)

    n = d.styles['Normal']
    n.font.name = FONT; n.font.size = Pt(10.5); n.font.color.rgb = INK
    n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.12

    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(0)
    r = t.add_run("LEAD DETAILS"); r.bold = True; r.font.size = Pt(16)
    r.font.name = FONT; r.font.color.rgb = INK

    sub = d.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run("Lingtu Law Office  ·  Case Referral")
    r.font.size = Pt(9); r.font.color.rgb = GREY; r.font.name = FONT
    _rule(sub)

    for heading, fields in SCHEMA:
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(5)
        r = h.add_run(heading.upper()); r.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = GREY; r.font.name = FONT
        r.font.all_caps = True
        _rule(h, 4)
        for label, key, indent in fields:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            if indent:
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.space_after = Pt(7)
            sep = "  " if label.endswith("?") else ":  "
            lr = p.add_run(f"{label}{sep}"); lr.bold = True; lr.font.name = FONT
            lr.font.size = Pt(10.5); lr.font.color.rgb = INK
            vr = p.add_run(str(data.get(key, f"[{label}]")))
            vr.font.name = FONT; vr.font.size = Pt(10.5); vr.font.color.rgb = INK

    d.save(out)
    return out


if __name__ == "__main__":
    out = sys.argv[1]
    data = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else BLANK
    print("saved", build(data, out))
