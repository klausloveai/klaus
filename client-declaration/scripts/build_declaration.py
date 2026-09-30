# -*- coding: utf-8 -*-
"""Build a single bilingual (EN/CN) client statement + signable declaration .docx.

Usage:  python3 build_declaration.py config.json

Config keys — see references/example-config.json for a filled one.
  out                 output .docx path
  title_en/title_cn   document title
  case_lines          list of centred grey subtitle lines (strings)
  narrative_head_en/_cn   e.g. "Statement of Peiyun Zhou" / "周培云陈述"
  timeline            [[time, english, 中文], ...]           (omit or [] to skip Part I)
  narrative           [[english, 中文], ...]                  (omit or [] to skip Part II)
  declarants          [{name_en, name_cn, body_en, body_cn}, ...]
  translator          true|false  (Certificate of Translation block)

**bold** inside any string renders bold.
"""
import json, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = '1F3864'
PERJURY_EN = ("I declare under penalty of perjury under the laws of the State of California "
              "that the foregoing is true and correct.")
PERJURY_CN = "本人依加利福尼亚州法律，在伪证罪处罚下声明，以上内容真实无误。"
EXEC_EN = "Executed on ______________________, {year}, at ______________________________, California."
EXEC_CN = "签署日期：{year} 年 ______ 月 ______ 日　　签署地点：加利福尼亚州 ______________________"


def build(cfg):
    d = Document(); s = d.sections[0]
    s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)
    st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    st.paragraph_format.space_after = Pt(0)
    year = cfg.get('year', 2026)

    def setfont(r, cn=False, size=10.5):
        r.font.size = Pt(size)
        rf = r._element.get_or_add_rPr().get_or_add_rFonts()
        if cn:
            for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                rf.set(qn(a), 'SimSun')
        else:
            rf.set(qn('w:ascii'), 'Times New Roman')
            rf.set(qn('w:hAnsi'), 'Times New Roman')
            rf.set(qn('w:eastAsia'), 'SimSun')

    def emit(p, txt, cn, size, bold, white=False, color=None, italic=False):
        for part in re.split(r'(\*\*.+?\*\*)', txt):
            if not part:
                continue
            b = part.startswith('**')
            r = p.add_run(part[2:-2] if b else part)
            r.bold = (b or bold); r.italic = italic
            setfont(r, cn, size)
            if white:
                r.font.color.rgb = RGBColor.from_string('FFFFFF')
            elif color:
                r.font.color.rgb = RGBColor.from_string(color)

    def P(txt='', size=10.5, bold=False, cn=False, after=4, before=0,
          align=None, color=None, italic=False):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'justify':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        emit(p, txt, cn, size, bold, color=color, italic=italic)
        return p

    def shade(c, h):
        e = OxmlElement('w:shd'); e.set(qn('w:fill'), h)
        c._tc.get_or_add_tcPr().append(e)

    def cell(c, txt, cn=False, size=9.5, bold=False, white=False):
        c.text = ''
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        emit(p, txt, cn, size, bold, white=white)

    def table(cols, rows):
        """cols = [(header, width_in, is_chinese), ...]"""
        t = d.add_table(rows=1, cols=len(cols)); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, (h, w, isc) in enumerate(cols):
            c = t.rows[0].cells[i]; c.width = Inches(w)
            cell(c, h, cn=isc, bold=True, white=True); shade(c, NAVY)
        for r in rows:
            row = t.add_row()
            for i, (h, w, isc) in enumerate(cols):
                row.cells[i].width = Inches(w)
                cell(row.cells[i], r[i], cn=isc, bold=(i == 0 and len(cols) == 3))
        return t

    # ---- header ----
    P(cfg['title_en'], size=15, bold=True, align='center', after=2)
    P(cfg['title_cn'], size=15, bold=True, cn=True, align='center', after=4)
    for i, ln in enumerate(cfg.get('case_lines', [])):
        P(ln, size=10 if i == 0 else 9.5, align='center', color='595959', after=1)
    P('', after=9)
    P("Please read each paragraph. The English and the Chinese say the same thing. "
      "If anything is wrong, or if anything is missing, tell us before you sign.",
      size=9.5, italic=True, after=1)
    P("请逐段阅读。左右两栏内容相同。如有任何不符或遗漏之处，请在签字前告诉我们。",
      size=9.5, italic=True, cn=True, after=10)

    # ---- I. timeline ----
    if cfg.get('timeline'):
        P("I.   TIMELINE  /  一、时间轴", size=12.5, bold=True, color=NAVY, before=4, after=4)
        table([("Time / 时间", 1.0, False), ("English", 3.0, False), ("中文", 3.0, True)],
              cfg['timeline'])
        d.add_page_break()

    # ---- II. narrative ----
    if cfg.get('narrative'):
        P("II.   STATEMENT  /  二、事实陈述", size=12.5, bold=True, color=NAVY, before=2, after=4)
        if cfg.get('narrative_head_en'):
            P(f"{cfg['narrative_head_en']}  /  {cfg.get('narrative_head_cn','')}",
              size=10, bold=True, after=4)
        table([("English", 3.5, False), ("中文", 3.5, True)], cfg['narrative'])
        d.add_page_break()

    # ---- III. declaration ----
    P("III.   DECLARATION  /  三、声明", size=12.5, bold=True, color=NAVY, before=2, after=6)

    def sigblock(name_en, name_cn):
        P(PERJURY_EN, size=10, bold=True, after=3)
        P(PERJURY_CN, size=10, bold=True, cn=True, after=10)
        P(EXEC_EN.format(year=year), size=10, after=3)
        P(EXEC_CN.format(year=year), size=10, cn=True, after=18)
        P("____________________________________________", size=10, after=2)
        if name_en:
            P(f"{name_en}  /  {name_cn}", size=10, bold=True, after=16)
        else:
            P("Name / 姓名: ______________________________", size=10, after=16)

    for dec in cfg.get('declarants', []):
        P(dec['body_en'], size=10, align='justify', after=3)
        P(dec['body_cn'], size=10, cn=True, align='justify', after=8)
        sigblock(dec['name_en'], dec['name_cn'])

    if cfg.get('translator', True):
        P("CERTIFICATE OF TRANSLATION  /  翻译人声明",
          size=11, bold=True, color=NAVY, before=6, after=5)
        P("I am fluent in English and in Mandarin Chinese. I translated the foregoing statement "
          "from English into Mandarin Chinese, and I read the Chinese text to the declarants in "
          "Mandarin Chinese before they signed. The translation is true and complete to the best "
          "of my ability.", size=10, align='justify', after=3)
        P("本人通晓英文与普通话。以上陈述由本人自英文译为中文，并于声明人签字前以普通话向其宣读中文文本。"
          "就本人能力所及，该翻译真实完整。", size=10, cn=True, align='justify', after=8)
        sigblock(None, None)

    d.save(cfg['out'])
    return cfg['out']


if __name__ == '__main__':
    cfg = json.load(open(sys.argv[1], encoding='utf-8'))
    print('saved', build(cfg))
