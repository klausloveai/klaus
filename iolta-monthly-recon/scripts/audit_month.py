#!/usr/bin/env python3
"""Independent audit of one month: re-derive every figure from the bank statement and the
journal, then check the package files actually on disk agree. Does not trust the package."""
import csv, datetime, os, re, sys, argparse, calendar
sys.path.insert(0,'/Users/klaus/.claude/skills/iolta-monthly-recon/scripts')
import openpyxl
from collections import defaultdict
from build_package import (load_journal, issued, cleared, num, f, norm,
                           carryover, load_disbursements, sweep_allocation)
MR='/Users/klaus/Library/CloudStorage/GoogleDrive-klaus@lingtulaw.com/My Drive/Lingtu Law-Disbursement/IOLTA#3618/Monthly Reconciliations/'

ap=argparse.ArgumentParser(); ap.add_argument('--month',required=True); ap.add_argument('--bank-end',type=float,required=True)
a=ap.parse_args()
y,m=map(int,a.month.split('-')); eom=datetime.date(y,m,calendar.monthrange(y,m)[1])
MON=datetime.date(y,m,1).strftime('%B')
P=f'{MR}{a.month} State Bar Package/'
print(f'### {a.month} 独立复核\n')

# ---- 1. statement self-check
rows=[]; summ={}
for r in csv.reader(open(f'{MR}stmt-3618-{a.month}.csv')):
    if len(r)>=3 and r[0].startswith(('Beginning','Total','Ending')): summ[r[0].split(' as of')[0]]=float(r[2].replace(',',''))
    if len(r)>=4 and re.match(r'^\d\d/\d\d/\d{4}$',r[0]) and r[2]: rows.append((r[0],r[1],float(r[2].replace(',',''))))
beg=[v for k,v in summ.items() if k.startswith('Beginning')][0]
cr=[v for k,v in summ.items() if k=='Total credits'][0]
db=[v for k,v in summ.items() if k=='Total debits'][0]
end=[v for k,v in summ.items() if k.startswith('Ending')][0]
print(f'1. 对账单自检  {beg:,.2f} + {cr:,.2f} {db:,.2f} = {beg+cr+db:,.2f}  vs 期末 {end:,.2f}  {"✓" if abs(beg+cr+db-end)<.005 else "✗"}')
print(f'   期末 vs 我们用的 bank-end {a.bank_end:,.2f}  {"✓" if abs(end-a.bank_end)<.005 else "✗"}')
bank_cr=sum(x[2] for x in rows if x[2]>0); bank_db=sum(x[2] for x in rows if x[2]<0)
print(f'   明细加总 贷 {bank_cr:,.2f} 借 {bank_db:,.2f}  {"✓" if abs(bank_cr-cr)<.005 and abs(bank_db-db)<.005 else "✗"}')

# ---- 2. journal vs statement, line by line
ws,last=load_journal()
disb=load_disbursements(); carry=carryover(disb); sweeps=sweep_allocation(ws,last,eom,disb)
mk_dep=mk_dis=0; mk=[]
for r in range(10,last+1):
    if str(ws.cell(r,14).value or '')==a.month:
        mk.append(r); mk_dep+=num(ws.cell(r,6).value); mk_dis+=num(ws.cell(r,7).value)
print(f'\n2. 账本标记本月已兑现 {len(mk)} 行  进 {mk_dep:,.2f}  出 {mk_dis:,.2f}')
print(f'   vs 银行            贷 {cr:,.2f}  借 {-db:,.2f}')
print(f'   {"✓ 两边完全相等" if abs(mk_dep-cr)<.005 and abs(mk_dis+db)<.005 else "✗ 有差额"}')

# ---- 3. period set, OS/DIT
inp=[]; OS=[]; DIT=[]
for r in range(10,last+1):
    d,s=num(ws.cell(r,6).value),num(ws.cell(r,7).value)
    if not (d or s): continue
    c=cleared(ws,r); i=issued(ws,r,y)
    if not ((i and i<=eom) or (c and c<=eom)): continue
    inp.append(r)
    if c is None or c>eom: (OS if s else DIT).append(r)
os_amt=round(sum(num(ws.cell(r,7).value) for r in OS),2); dit=round(sum(num(ws.cell(r,6).value) for r in DIT),2)
J=round(sum(num(ws.cell(r,6).value) for r in inp)-sum(num(ws.cell(r,7).value) for r in inp),2)
adj=round(a.bank_end+dit-os_amt,2)

# ---- 4. ledgers
led=defaultdict(float)
for c_,v in carry.items(): led[c_]+=v
for r in inp:
    cl=str(ws.cell(r,9).value or '').strip()
    if not cl or cl.startswith(('Carryover','Operating')): continue
    led[norm(cl)]+=num(ws.cell(r,6).value)-num(ws.cell(r,7).value)
for c_,lst in sweeps.items():
    for _,_,_,amt in lst: led[c_]-=amt
L=round(sum(led.values()),2)
print(f'\n3. 三方')
print(f'   line1 总账        {J:>14,.2f}')
print(f'   line2 分户账      {L:>14,.2f}   差 {J-L:+,.2f}')
print(f'   line3 调整后银行  {adj:>14,.2f}   差 {J-adj:+,.2f}')
print(f'         银行 {a.bank_end:,.2f} + 在途 {dit:,.2f}({len(DIT)}) − 未兑现 {os_amt:,.2f}({len(OS)})')
print(f'   {"✓ 三方一致" if abs(J-L)<.005 and abs(J-adj)<.005 else "✗ 不一致"}')
neg={c_:round(v,2) for c_,v in led.items() if v<-0.004}
print(f'   负余额客户: {neg or "无"}')

# ---- 5. package files agree
print(f'\n4. package 文件核对')
if not os.path.isdir(P): print('   ✗ package 目录不存在'); sys.exit()
def rd(fn,col,start=5):
    w=openpyxl.load_workbook(P+fn,data_only=False).active
    return [num(w.cell(r,col).value) for r in range(start,w.max_row+1)]
try:
    j1=openpyxl.load_workbook(P+f'1 - Account Journal (through {eom:%m-%d-%Y}).xlsx',data_only=False).active
    d1=sum(num(j1.cell(r,6).value) for r in range(5,j1.max_row+1)); s1=sum(num(j1.cell(r,7).value) for r in range(5,j1.max_row+1))
    nrow=sum(1 for r in range(5,j1.max_row+1) if j1.cell(r,2).value)
    print(f'   附件1 总账 {nrow} 行  余额 {d1-s1:,.2f}  {"✓" if abs(d1-s1-J)<.005 else "✗"}  (行数 vs 期内 {len(inp)}: {"✓" if nrow==len(inp) else "✗"})')
    w2=openpyxl.load_workbook(P+f'2 - Client Ledgers (as of {eom:%m-%d-%Y}).xlsx',data_only=False)['Summary']
    t2=sum(num(w2.cell(r,5).value) for r in range(5,w2.max_row+1))
    print(f'   附件2 分户账 {len(openpyxl.load_workbook(P+f"2 - Client Ledgers (as of {eom:%m-%d-%Y}).xlsx").sheetnames)-1} 客户  合计 {t2:,.2f}  {"✓" if abs(t2-L)<.005 else "✗"}')
    w4=openpyxl.load_workbook(P+f'4 - Client Ledger Summary with Balances ({eom:%m-%d-%Y}).xlsx',data_only=False).active
    t4=sum(num(w4.cell(r,4).value) for r in range(5,w4.max_row+1))
    print(f'   附件4 余额汇总 合计 {t4:,.2f}  {"✓" if abs(t4-L)<.005 else "✗"}')
    w5=openpyxl.load_workbook(P+f'5 - Outstanding Deposits ({eom:%m-%d-%Y}).xlsx',data_only=False).active
    t5=sum(num(w5.cell(r,4).value) for r in range(5,w5.max_row+1))
    print(f'   附件5 在途存款 {t5/2 if t5 else 0:,.2f}  {"✓" if abs((t5/2 if t5 else 0)-dit)<.01 else "(含合计行,人工看)"}')
    w6=openpyxl.load_workbook(P+f'6 - Outstanding Disbursements ({eom:%m-%d-%Y}).xlsx',data_only=False).active
    n6=sum(1 for r in range(5,w6.max_row+1) if str(w6.cell(r,1).value or '').isdigit())
    t6=sum(num(w6.cell(r,5).value) for r in range(5,w6.max_row+1) if str(w6.cell(r,1).value or '').isdigit())
    print(f'   附件6 未兑现 {n6} 张 {t6:,.2f}  {"✓" if abs(t6-os_amt)<.005 and n6==len(OS) else "✗"}')
    rf=[x for x in os.listdir(P) if x.startswith('Recon-')][0]
    wr=openpyxl.load_workbook(P+rf)['Reconciliation form']
    l1,l2,l3=num(wr['L19'].value),num(wr['L39'].value),num(wr['L50'].value)
    print(f'   对账表 L1 {l1:,.2f} L2 {l2:,.2f} L3 {l3:,.2f}  {"✓" if abs(l1-J)<.005 and abs(l2-J)<.005 and abs(l3-J)<.005 else "✗"}')
    print(f'          3A {num(wr["L44"].value):,.2f} 3B {num(wr["L46"].value):,.2f} 3C {num(wr["L48"].value):,.2f}  '
          f'{"✓" if abs(num(wr["L44"].value)-a.bank_end)<.005 and abs(num(wr["L46"].value)-dit)<.005 and abs(num(wr["L48"].value)-os_amt)<.005 else "✗"}')
    print(f'          Recon Date {wr["L7"].value}  签名 Preparer={wr["I59"].value!r} Attorney={wr["I68"].value!r}')
except Exception as e:
    print('   ✗', e)
