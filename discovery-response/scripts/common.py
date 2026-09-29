# -*- coding: utf-8 -*-
"""Pleading-paper docx engine + the objection library.

The base document (assets/pleading-paper-base.docx) carries its fonts on the RUNS,
not on the styles, and its 1-28 line numbers live in header1.xml as real text. So
every new paragraph is CLONED from a prototype paragraph in the base and re-texted.
Wiping runs and calling add_run() silently drops the whole document to Courier, and
leaving a stray <w:tab/> in a cloned run shoves the line to the paragraph's tab stop.
Only the pleading chrome is reused from that file -- never its content or its
objection style (see references/hernan-standard.md).
"""
import copy, re, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH

BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "assets", "pleading-paper-base.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# ------------------------------------------------------- objection library
# One ground, one cite, no stacking. Calibrated to Hernán's signed work
# (Zhiping Liu v. State Farm: 19 objections in 63 answers).
def _o(body, k="interrogatory"):
    return "Responding Party objects to this {k} to the extent that it {b}".format(k=k, b=body)

O_WP        = lambda k="interrogatory": _o("seeks material protected by the attorney work product doctrine under Code of Civil Procedure section 2018.030.", k)
O_PRIVACY   = lambda k="interrogatory": _o("seeks information within the zone of privacy protected by article I, section 1 of the California Constitution.", k)
O_EXPERT    = lambda k="interrogatory": _o("calls for expert opinion in advance of the disclosure deadlines set by Code of Civil Procedure sections 2034.210 through 2034.310.", k)
O_BRITT     = lambda k="interrogatory": _o("seeks information within the zone of privacy protected by article I, section 1 of the California Constitution and is not limited to the condition placed at issue in this action. (Britt v. Superior Court (1978) 20 Cal.3d 844, 863-864.)", k)
O_PREMATURE = lambda k="interrogatory": _o("calls for a legal conclusion and seeks contentions that are premature at this stage of the litigation.", k)

SWO1 = "Subject to and without waiving that objection, Responding Party responds:"
SWO2 = "Subject to and without waiving those objections, Responding Party responds:"

# Mandatory whenever an objection is made but everything is still produced.
NO_WITHHOLD = "No responsive document is being withheld on the basis of the objection stated above."

CONT   = ("Discovery and investigation are continuing, and Responding Party reserves the right to "
          "amend and to supplement this response under Code of Civil Procedure section 2030.310.")
CONT_P = ("Discovery and investigation are continuing, and Responding Party reserves the right to "
          "supplement this production as further documents are received.")

COMPLY = ("Responding Party will comply with this request in whole. All documents in Responding "
          "Party's possession, custody or control to which no objection is being made are produced "
          "concurrently with this response.")
UNABLE = ("Responding Party is unable to comply with this request because no responsive document has "
          "ever existed, or ever existed in Responding Party's possession, custody or control. "
          "Responding Party has conducted a diligent search and a reasonable inquiry.")

PRELIM = ("These answers are solely for the purpose of, and in relation to, this action. Each answer is "
          "given subject to all appropriate objections (including, but not limited to, objections covering "
          "competency, relevancy, materiality, propriety, and admissibility) which would require the "
          "exclusion of any statement contained herein if it were made by a witness present and testifying "
          "in court. All such objections and grounds therefore are reserved and may be ruled on at the time "
          "of trial. The party on whose behalf these answers are given has not yet completed its "
          "investigation of the facts relating to this action, has not yet completed discovery in this "
          "action, and has not yet completed preparation for trial. Consequently, the following answers are "
          "given without prejudice to the answering party's right to produce at the time of trial "
          "subsequently-discovered evidence relating to the proof of facts subsequently discovered to be "
          "material.")

# Verification for a client who does not read English with ease -- required whenever
# FROG 2.9 / 2.10 are answered "no". A "I have read the foregoing" verification
# contradicts those answers.
VERIFICATION_TRANSLATED = (
    "I am the Plaintiff in the above-entitled action. My native language is Mandarin Chinese, and I "
    "do not read English with ease. The foregoing was translated to me from English into Mandarin "
    "Chinese, and I know its contents. The same is true of my own knowledge, except as to those "
    "matters which are stated on information and belief, and as to those matters I believe them to "
    "be true.")
TRANSLATOR_DECL = (
    "I am fluent in English and in Mandarin Chinese. On ______________, 2026, I translated the "
    "foregoing responses and the verification from English into Mandarin Chinese for YI CONG, who "
    "stated that he understood them. I declare under penalty of perjury under the laws of the State "
    "of California that the foregoing is true and correct.")

# ----------------------------------------------------------------- docx glue
HL = re.compile(r"«(.+?)»", re.S)      # «...» -> yellow highlight

def _runs(p):
    return p._p.findall(W + "r")

def retext(p, text, bold=None):
    """Replace a paragraph's text, keeping run 0's character formatting.

    A literal \\t becomes a real <w:tab/>. «spans» are cloned from run 0 and
    highlighted yellow. Every w:t / w:tab / w:br already in the prototype run is
    cleared first.
    """
    rs = _runs(p)
    if not rs:
        r = p.add_run(text)
        if bold is not None:
            r.bold = bold
        return
    proto = rs[0]
    for r in rs[1:]:
        r.getparent().remove(r)
    for tag in ("t", "tab", "br"):
        for el in proto.findall(W + tag):
            proto.remove(el)

    segs, pos = [], 0
    for m in HL.finditer(text):
        if m.start() > pos:
            segs.append((text[pos:m.start()], False))
        segs.append((m.group(1), True))
        pos = m.end()
    if pos < len(text) or not segs:
        segs.append((text[pos:], False))

    def fill(r, seg):
        for chunk in re.split(r"(\t)", seg):
            if chunk == "\t":
                r.append(r.makeelement(W + "tab", {}))
            elif chunk:
                t = r.makeelement(W + "t", {})
                t.text = chunk
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)

    def style(r, hl):
        if not hl and bold is None:
            return
        rPr = r.find(W + "rPr")
        if rPr is None:
            rPr = r.makeelement(W + "rPr", {}); r.insert(0, rPr)
        if hl:
            for h in rPr.findall(W + "highlight"):
                rPr.remove(h)
            rPr.append(rPr.makeelement(W + "highlight", {W + "val": "yellow"}))
        if bold is not None:
            for b in rPr.findall(W + "b"):
                rPr.remove(b)
            if bold:
                rPr.insert(0, rPr.makeelement(W + "b", {}))

    anchor = proto
    for i, (seg, hl) in enumerate(segs):
        if i == 0:
            fill(proto, seg); style(proto, hl)
        else:
            r = copy.deepcopy(proto)
            for tag in ("t", "tab", "br"):
                for el in r.findall(W + tag):
                    r.remove(el)
            fill(r, seg); style(r, hl)
            anchor.addnext(r); anchor = r


def set_cell(cell, lines):
    """Rewrite a caption cell; each line is cloned from its first paragraph."""
    proto = copy.deepcopy(cell.paragraphs[0]._p)
    for p in list(cell.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for _ in lines[1:]:
        cell._tc.append(copy.deepcopy(proto))
    for para_, (txt, bold, align) in zip(cell.paragraphs, lines):
        retext(para_, txt, bold=bold)
        if align is not None:
            para_.alignment = align


def load_base():
    return Document(BASE)


def _plain_sub(p_el):
    """The base's lettered sub-answer is an auto-numbered list item; strip the
    numbering and give it a hanging indent so "(a) ..." lines up and wraps."""
    el = copy.deepcopy(p_el)
    pPr = el.find(W + "pPr")
    for n in pPr.findall(W + "numPr"):
        pPr.remove(n)
    for i in pPr.findall(W + "ind"):
        pPr.remove(i)
    ind = pPr.makeelement(W + "ind", {W + "start": "1080", W + "hanging": "360",
                                      W + "left": "1080", W + "firstLine": "0"})
    pPr.insert(len(pPr) - 1, ind)
    return el


def grab_protos(doc):
    """Prototype paragraphs, captured BEFORE the body is truncated."""
    ps = doc.paragraphs
    return {
        "head":   copy.deepcopy(ps[19]._p),   # Heading 1, bold centered
        "prelim": copy.deepcopy(ps[20]._p),   # justified body block
        "label":  copy.deepcopy(ps[22]._p),   # "RESPONSE TO ... NO. x:"
        "body":   copy.deepcopy(ps[23]._p),   # normal, tab-indented first line
        "sub":    _plain_sub(ps[25]._p),      # lettered sub-answer
        "plain":  copy.deepcopy(ps[16]._p),   # PROPOUNDING PARTY line
    }


def rewrite_head(doc, atty_lines, court_lines, left_cell, right_cell, party_lines):
    ps = doc.paragraphs
    for i in range(0, 11):
        retext(ps[i], atty_lines[i] if i < len(atty_lines) else "")
    for i, txt in zip((13, 14), court_lines):
        retext(ps[i], txt)
    t = doc.tables[0]
    set_cell(t.rows[0].cells[0], left_cell)
    set_cell(t.rows[0].cells[1], right_cell)
    for i, txt in zip((16, 17, 18), party_lines):
        retext(ps[i], txt)


def truncate_body(doc, keep_upto=18):
    for p in list(doc.paragraphs)[keep_upto + 1:]:
        p._p.getparent().remove(p._p)


def emit(doc, protos, kind, text, bold=None):
    """Append a clone of prototype `kind` carrying `text`, before the sectPr."""
    from docx.text.paragraph import Paragraph
    new = copy.deepcopy(protos[kind])
    body = doc.element.body
    sect = body.find(W + "sectPr")
    (sect.addprevious(new) if sect is not None else body.append(new))
    p = Paragraph(new, doc._body)
    retext(p, text, bold=bold)
    return p


def page_break(doc, protos):
    p = emit(doc, protos, "body", "")
    r = _runs(p)[0]
    r.insert(0, r.makeelement(W + "br", {W + "type": "page"}))
    return p


def set_footer(doc, text):
    """Footer paragraph 1 carries the title; paragraph 0 is the PAGE field."""
    retext(doc.sections[0].footer.paragraphs[1], text)
